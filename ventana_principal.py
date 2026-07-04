import traceback
import psycopg2
from datetime import datetime, timedelta

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, 
    QLabel, QLineEdit, QPushButton, QComboBox, QTableWidget, 
    QTableWidgetItem, QStackedWidget, QFormLayout, QHeaderView, 
    QMessageBox, QDialog, QGridLayout, QFrame, QSizePolicy, QCheckBox
)
from PyQt6.QtCore import Qt, QRegularExpression
from PyQt6.QtGui import QPixmap, QRegularExpressionValidator, QValidator
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import base_datos

# =====================================================================
# DIÁLOGO DE EDICIÓN DE USUARIO
# =====================================================================
class DialogoEditarUsuario(QDialog):
    def __init__(self, datos, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Editar Usuario")
        self.setMinimumWidth(400)
        self.id_usuario = datos[0]
        
        layout = QVBoxLayout(self)
        form = QFormLayout()
        
        self.txt_cedula = QLineEdit(datos[1])
        self.txt_matricula = QLineEdit(datos[2])
        self.txt_nombre = QLineEdit(datos[3])
        
        self.cmb_tipo = QComboBox()
        self.cmb_tipo.addItems(["Estudiante", "Profesor", "Empleado", "Invitado"])
        self.cmb_tipo.setCurrentText(datos[4])
        
        self.cmb_carrera = QComboBox()
        carreras = ["", "Ninguna / No Aplica", "Administración de Empresas", "Ingeniería en Sistemas Computacionales", "Arquitectura", "Derecho", "Medicina", "Psicología"]
        self.cmb_carrera.addItems(carreras)
        self.cmb_carrera.setEditable(True)
        self.cmb_carrera.setCurrentText(datos[5] if datos[5] else "")
        
        form.addRow("Cédula:", self.txt_cedula)
        form.addRow("Matrícula:", self.txt_matricula)
        form.addRow("Nombre:", self.txt_nombre)
        form.addRow("Tipo:", self.cmb_tipo)
        form.addRow("Carrera:", self.cmb_carrera)
        
        layout.addLayout(form)
        
        btn_guardar = QPushButton("Guardar Cambios")
        btn_guardar.setStyleSheet("background-color: #10B981; color: white; font-weight: bold; padding: 10px;")
        btn_guardar.clicked.connect(self.accept)
        layout.addWidget(btn_guardar)

    def obtener_datos(self):
        return (
            self.txt_cedula.text().strip(),
            self.txt_matricula.text().strip(),
            self.txt_nombre.text().strip(),
            self.cmb_tipo.currentText(),
            self.cmb_carrera.currentText()
        )

# =====================================================================
# 3. INTERFAZ PRINCIPAL
# =====================================================================
class SistemaBiblioteca(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("UTESA Plus - Control de Acceso Salón de Lectura")
        self.resize(1300, 750)
        
        self.usuario_actual_prestamo = None
        self.stacked_widget = QStackedWidget()
        self.setCentralWidget(self.stacked_widget)
        
        self.vista_formulario_principal = QWidget()
        self.vista_kiosco_acceso = QWidget()
        self.vista_metricas = QWidget()
        self.vista_usuarios = QWidget()
        
        self.stacked_widget.addWidget(self.vista_formulario_principal)
        self.stacked_widget.addWidget(self.vista_kiosco_acceso)
        self.stacked_widget.addWidget(self.vista_metricas)
        self.stacked_widget.addWidget(self.vista_usuarios)
        
        self.crear_vista_formulario_principal()
        self.crear_vista_kiosco_acceso()
        self.crear_vista_metricas()
        self.crear_vista_usuarios()
        
        self.stacked_widget.setCurrentIndex(0)

    def cargar_imagen_branding(self, size=80):
        lbl_img = QLabel()
        pixmap = QPixmap("mi_foto.png")
        if not pixmap.isNull():
            lbl_img.setPixmap(pixmap.scaled(size, size, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
        else:
            lbl_img.setText("👤")
            lbl_img.setStyleSheet(f"font-size: {size//2}px; color: #94A3B8;")
        lbl_img.setAlignment(Qt.AlignmentFlag.AlignCenter)
        return lbl_img

    def barra_navegacion(self, index_actual):
        layout = QHBoxLayout()
        layout.setContentsMargins(10, 10, 10, 10)
        
        btn1 = QPushButton("Formulario de Asignación de PCs")
        btn2 = QPushButton("Acceso Táctil")
        btn3 = QPushButton("Métricas y Reportes")
        btn4 = QPushButton("Gestión de Usuarios")
        
        for btn in [btn1, btn2, btn3, btn4]:
            btn.setMinimumHeight(40)
            btn.setStyleSheet("background-color: #E2E8F0; color: #1E293B; font-weight: 500;")
            
        btn1.clicked.connect(lambda: self.stacked_widget.setCurrentIndex(0))
        btn2.clicked.connect(self.activar_modo_kiosco)
        btn3.clicked.connect(lambda: self.stacked_widget.setCurrentIndex(2))
        btn4.clicked.connect(lambda: self.stacked_widget.setCurrentIndex(3))
        
        if index_actual == 0: btn1.setStyleSheet("background-color: #0EA5E9; color: white; font-weight: bold;")
        if index_actual == 1: btn2.setStyleSheet("background-color: #0EA5E9; color: white; font-weight: bold;")
        if index_actual == 2: btn3.setStyleSheet("background-color: #0EA5E9; color: white; font-weight: bold;")
        if index_actual == 3: btn4.setStyleSheet("background-color: #0EA5E9; color: white; font-weight: bold;")
        
        layout.addWidget(btn1)
        layout.addWidget(btn2)
        layout.addWidget(btn3)
        layout.addWidget(btn4)
        return layout

    # =====================================================================
    # VISTA 1: ASIGNACIÓN PC
    # =====================================================================
    def crear_vista_formulario_principal(self):
        layout_global = QVBoxLayout(self.vista_formulario_principal)
        layout_global.addLayout(self.barra_navegacion(0))
        layout_contenido = QHBoxLayout()
        
        sec_formulario = QVBoxLayout()
        lbl_f_titulo = QLabel("Asignación de Computadoras")
        lbl_f_titulo.setStyleSheet("font-size: 18px; font-weight: bold; color: #0F172A; margin-bottom: 10px;")
        sec_formulario.addWidget(lbl_f_titulo)
        
        form_layout = QFormLayout()
        self.txt_form_buscar = QLineEdit()
        self.txt_form_buscar.setPlaceholderText("Ingrese Cédula o Matrícula y presione Enter")
        self.txt_form_buscar.setMinimumHeight(35)
        self.txt_form_buscar.returnPressed.connect(self.buscar_usuario_formulario)
        
        self.lbl_form_datos = QLabel("Usuario: No seleccionado\nTipo: -\nCarrera: -")
        self.lbl_form_datos.setStyleSheet("background-color: #F1F5F9; color: #1E293B; padding: 10px; border-radius: 6px; font-size: 14px;")
        
        self.txt_form_fecha = QLineEdit(datetime.now().strftime("%Y-%m-%d"))
        self.txt_form_fecha.setReadOnly(True)
        self.txt_form_fecha.setMinimumHeight(35)
        
        self.cmb_form_pcs = QComboBox()
        self.cmb_form_pcs.setMinimumHeight(35)
        self.actualizar_combo_pcs()
        
        form_layout.addRow("Identificación:", self.txt_form_buscar)
        form_layout.addRow("Datos del Usuario:", self.lbl_form_datos)
        form_layout.addRow("Día (Actual):", self.txt_form_fecha)
        form_layout.addRow("PC a Asignar:", self.cmb_form_pcs)
        sec_formulario.addLayout(form_layout)
        
        btn_guardar_p = QPushButton("Guardar Asignación")
        btn_guardar_p.setStyleSheet("background-color: #10B981; color: white; font-weight: bold; min-height: 45px; font-size: 15px;")
        btn_guardar_p.clicked.connect(self.guardar_asignacion_pc)
        sec_formulario.addWidget(btn_guardar_p)
        sec_formulario.addStretch()
        
        sec_edicion = QVBoxLayout()
        lbl_e_titulo = QLabel("Editor de Formulario (PCs)")
        lbl_e_titulo.setStyleSheet("font-size: 18px; font-weight: bold; color: #0F172A; margin-bottom: 10px;")
        sec_edicion.addWidget(lbl_e_titulo)
        
        self.table_pcs = QTableWidget(0, 2)
        self.table_pcs.setHorizontalHeaderLabels(["Nombre PC", "Estado"])
        self.table_pcs.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.actualizar_tabla_pcs()
        sec_edicion.addWidget(self.table_pcs)
        
        layout_agregar_pc = QHBoxLayout()
        self.txt_nueva_pc = QLineEdit()
        self.txt_nueva_pc.setPlaceholderText("Ej: PC 5")
        self.txt_nueva_pc.setMinimumHeight(35)
        btn_agregar_pc = QPushButton("Agregar PC")
        btn_agregar_pc.setStyleSheet("background-color: #0F172A; color: white; font-weight: 500; min-height: 35px;")
        btn_agregar_pc.clicked.connect(self.agregar_nueva_pc)
        
        layout_agregar_pc.addWidget(self.txt_nueva_pc)
        layout_agregar_pc.addWidget(btn_agregar_pc)
        sec_edicion.addLayout(layout_agregar_pc)
        
        btn_cambiar_estado = QPushButton("Alternar Activo / Inactivo")
        btn_cambiar_estado.setStyleSheet("background-color: #EF4444; color: white; min-height: 35px;")
        btn_cambiar_estado.clicked.connect(self.alternar_estado_pc)
        sec_edicion.addWidget(btn_cambiar_estado)
        
        frame_izq = QFrame()
        frame_izq.setLayout(sec_formulario)
        frame_der = QFrame()
        frame_der.setLayout(sec_edicion)
        
        layout_contenido.addWidget(frame_izq, 1)
        layout_contenido.addWidget(frame_der, 1)
        layout_global.addLayout(layout_contenido)

    def actualizar_combo_pcs(self):
        self.cmb_form_pcs.clear()
        equipos = base_datos.obtener_pcs_activas()
        for row in equipos:
            self.cmb_form_pcs.addItem(row[1], row[0])

    def actualizar_tabla_pcs(self):
        self.table_pcs.setRowCount(0)
        equipos = base_datos.obtener_todas_pcs()
        for row_idx, row in enumerate(equipos):
            self.table_pcs.insertRow(row_idx)
            item_nombre = QTableWidgetItem(row[1])
            item_nombre.setData(Qt.ItemDataRole.UserRole, row[0])
            item_estado = QTableWidgetItem("Activo" if row[2] == 1 else "Inactivo")
            item_nombre.setForeground(Qt.GlobalColor.black)
            item_estado.setForeground(Qt.GlobalColor.black)
            self.table_pcs.setItem(row_idx, 0, item_nombre)
            self.table_pcs.setItem(row_idx, 1, item_estado)

    def buscar_usuario_formulario(self):
        ident = self.txt_form_buscar.text().strip()
        if not ident: return
        res = base_datos.buscar_usuario(ident)
        if res:
            esta_bloqueado = False
            if len(res) > 4:
                esta_bloqueado = res[4]
            
            if esta_bloqueado:
                QMessageBox.warning(self, "Bloqueado", "Este usuario tiene el acceso bloqueado y no se le puede asignar PC.")
                return

            self.usuario_actual_prestamo = res[0]
            self.lbl_form_datos.setText(f"Usuario: {res[1]}\nTipo: {res[2]}\nCarrera: {res[3] if res[3] else 'N/A'}")
        else:
            self.usuario_actual_prestamo = None
            self.lbl_form_datos.setText("Usuario no registrado.")
            QMessageBox.warning(self, "Aviso", "El usuario no existe. Debe registrarse en el sistema.")

    def agregar_nueva_pc(self):
        nombre = self.txt_nueva_pc.text().strip()
        if not nombre: return
        try:
            base_datos.agregar_equipo(nombre)
            self.txt_nueva_pc.clear()
            self.actualizar_tabla_pcs()
            self.actualizar_combo_pcs()
        except psycopg2.IntegrityError:
            QMessageBox.critical(self, "Error", "Ese nombre de PC ya existe.")

    def alternar_estado_pc(self):
        fila_sel = self.table_pcs.currentRow()
        if fila_sel < 0: return
        id_pc = self.table_pcs.item(fila_sel, 0).data(Qt.ItemDataRole.UserRole)
        estado_actual = self.table_pcs.item(fila_sel, 1).text()
        nuevo_estado = 0 if estado_actual == "Activo" else 1
        base_datos.actualizar_estado_equipo(id_pc, nuevo_estado)
        self.actualizar_tabla_pcs()
        self.actualizar_combo_pcs()

    def guardar_asignacion_pc(self):
        if not self.usuario_actual_prestamo: return
        id_pc = self.cmb_form_pcs.currentData()
        if not id_pc: return
        base_datos.guardar_asignacion(self.usuario_actual_prestamo, id_pc)
        QMessageBox.information(self, "Éxito", "Asignación guardada correctamente.")
        self.txt_form_buscar.clear()
        self.lbl_form_datos.setText("Usuario: No seleccionado\nTipo: -\nCarrera: -")
        self.usuario_actual_prestamo = None

    # -----------------------------------------------------------------
    # VISTA 2: KIOSCO TÁCTIL
    # -----------------------------------------------------------------
    def crear_vista_kiosco_acceso(self):
        layout_global = QVBoxLayout(self.vista_kiosco_acceso)
        
        self.nav_widget_kiosco = QWidget()
        self.nav_widget_kiosco.setLayout(self.barra_navegacion(1))
        layout_global.addWidget(self.nav_widget_kiosco)
        
        # --- Selector de Recinto de Biblioteca ---
        self.lay_kiosco_top = QHBoxLayout()
        lbl_bib = QLabel("📍 Recinto Actual:")
        lbl_bib.setStyleSheet("font-weight: bold; font-size: 16px; color: #1E293B;")
        
        self.cmb_kiosco_biblioteca = QComboBox()
        self.cmb_kiosco_biblioteca.addItems([
            "Seleccione un recinto...",
            "Biblioteca Coral Mall",
            "Biblioteca Herrera H-A",
            "Biblioteca Herrera H-B",
            "Biblioteca Maximo Gomez Edif. C",
            "Biblioteca Maximo Gomez Edif. D",
            "Biblioteca Ozama"
        ])
        self.cmb_kiosco_biblioteca.setStyleSheet("font-size: 16px; padding: 5px; min-height: 40px;")
        self.cmb_kiosco_biblioteca.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        
        self.index_biblioteca_anterior = 0
        self.cmb_kiosco_biblioteca.currentIndexChanged.connect(self.verificar_cambio_biblioteca)
        
        self.lay_kiosco_top.addWidget(lbl_bib)
        self.lay_kiosco_top.addWidget(self.cmb_kiosco_biblioteca)
        self.lay_kiosco_top.addStretch(2)
        
        layout_global.addLayout(self.lay_kiosco_top)
        
        layout_kiosco_master = QVBoxLayout()
        layout_kiosco_master.setContentsMargins(20, 20, 20, 20)
        self.kiosco_stack = QStackedWidget()
        layout_kiosco_master.addWidget(self.kiosco_stack)
        
        # 2.1 INICIO KIOSCO
        self.p_kiosco_inicio = QWidget()
        lay_ini = QVBoxLayout(self.p_kiosco_inicio)
        lay_ini.setSpacing(20)
        
        lbl_k_tit = QLabel("SISTEMA DE ACCESO - SALÓN DE LECTURA")
        lbl_k_tit.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_k_tit.setStyleSheet("font-size: 32px; font-weight: bold; color: #0F172A; margin-bottom: 20px;")
        lay_ini.addWidget(lbl_k_tit)
        
        btn_k_ingresar = QPushButton("INGRESAR AL SALÓN")
        btn_k_ingresar.setStyleSheet("background-color: #0EA5E9; color: white; font-size: 20px; font-weight: bold; border-radius: 12px;")
        btn_k_ingresar.setFixedSize(500, 90)
        btn_k_ingresar.clicked.connect(lambda: self.kiosco_stack.setCurrentIndex(1))
        
        btn_k_registrar = QPushButton("REGISTRARME COMO NUEVO USUARIO")
        btn_k_registrar.setStyleSheet("background-color: #0F172A; color: white; font-size: 20px; font-weight: bold; border-radius: 12px;")
        btn_k_registrar.setFixedSize(500, 90)
        btn_k_registrar.clicked.connect(lambda: self.kiosco_stack.setCurrentIndex(2))
        
        self.btn_k_fullscreen = QPushButton("🖵 BLOQUEAR EN PANTALLA COMPLETA")
        self.btn_k_fullscreen.setStyleSheet("background-color: #1E293B; color: white; font-size: 14px; font-weight: bold; border-radius: 8px; padding: 10px;")
        self.btn_k_fullscreen.setFixedSize(350, 45)
        self.btn_k_fullscreen.clicked.connect(self.alternar_pantalla_completa)
        
        lay_ini.addStretch(1)
        lay_ini.addWidget(lbl_k_tit, 0, Qt.AlignmentFlag.AlignCenter)
        lay_ini.addWidget(btn_k_ingresar, 0, Qt.AlignmentFlag.AlignCenter)
        lay_ini.addWidget(btn_k_registrar, 0, Qt.AlignmentFlag.AlignCenter)
        lay_ini.addSpacing(30)
        lay_ini.addWidget(self.btn_k_fullscreen, 0, Qt.AlignmentFlag.AlignCenter)
        lay_ini.addStretch(1)
        
        self.kiosco_stack.addWidget(self.p_kiosco_inicio)
        
        # =====================================================================
        # 2.2 KEYPAD DE INGRESO
        # =====================================================================
        self.p_kiosco_ingreso = QWidget()
        lay_ing = QVBoxLayout(self.p_kiosco_ingreso) 
        lay_ing.setAlignment(Qt.AlignmentFlag.AlignCenter)

        form_ingreso_layout = QVBoxLayout()
        form_ingreso_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

        lbl_ing_tit = QLabel("Digite su Cédula o Matrícula")
        lbl_ing_tit.setStyleSheet("font-size: 24px; font-weight: bold; color: #1E293B;")
        lbl_ing_tit.setAlignment(Qt.AlignmentFlag.AlignCenter)

        self.txt_ingreso_id = QLineEdit()
        self.txt_ingreso_id.setStyleSheet("font-size: 28px; min-height: 60px; border: 3px solid #0EA5E9; border-radius: 8px; padding: 5px; background-color: white; color: #1E293B;")
        self.txt_ingreso_id.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.txt_ingreso_id.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)

        btn_ing_aceptar = QPushButton("CONFIRMAR ENTRADA")
        btn_ing_aceptar.setStyleSheet("background-color: #10B981; color: white; font-size: 20px; font-weight: bold; min-height: 60px; border-radius: 8px;")
        btn_ing_aceptar.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        btn_ing_aceptar.clicked.connect(self.procesar_ingreso_salon)

        h_form_ing = QHBoxLayout()
        h_form_ing.addStretch(1)
        v_form_inner = QVBoxLayout()
        v_form_inner.addWidget(lbl_ing_tit, 0, Qt.AlignmentFlag.AlignCenter)
        v_form_inner.addWidget(self.txt_ingreso_id)
        v_form_inner.addWidget(btn_ing_aceptar)
        h_form_ing.addLayout(v_form_inner, 3) 
        h_form_ing.addStretch(1)

        grid_keypad = QGridLayout()
        grid_keypad.setSpacing(10)
        grid_keypad.setAlignment(Qt.AlignmentFlag.AlignCenter)
        botones_pad = ['1','2','3','4','5','6','7','8','9','Borrar','0','Limpiar']
        row, col = 0, 0
        for boton in botones_pad:
            btn_pad = QPushButton(boton)
            btn_pad.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            btn_pad.setMinimumSize(90, 90)
            btn_pad.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
            btn_pad.setStyleSheet("background-color: white; color: #1E293B; border: 2px solid #CBD5E1; font-size: 22px; font-weight: bold; border-radius: 10px;")
            btn_pad.clicked.connect(self.tecla_num_presionada)
            grid_keypad.addWidget(btn_pad, row, col)
            col += 1
            if col > 2:
                col = 0
                row += 1
                
        h_pad_ing = QHBoxLayout()
        h_pad_ing.addStretch(1)
        h_pad_ing.addLayout(grid_keypad, 2)
        h_pad_ing.addStretch(1)

        btn_ing_volver = QPushButton("Volver Atrás")
        btn_ing_volver.setStyleSheet("background-color: #64748B; color: white; font-size: 16px; min-height: 45px;")
        btn_ing_volver.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        btn_ing_volver.clicked.connect(self.limpiar_volver_kiosco)
        
        h_volver_ing = QHBoxLayout()
        h_volver_ing.addStretch(1)
        h_volver_ing.addWidget(btn_ing_volver, 2)
        h_volver_ing.addStretch(1)

        lay_ing.addStretch(1)
        lay_ing.addLayout(h_form_ing)
        lay_ing.addSpacing(30)
        lay_ing.addLayout(h_pad_ing)
        lay_ing.addSpacing(30)
        lay_ing.addLayout(h_volver_ing)
        lay_ing.addStretch(1)

        self.kiosco_stack.addWidget(self.p_kiosco_ingreso)
        
        # =====================================================================
        # 2.3 REGISTRO TÁCTIL
        # =====================================================================
        self.p_kiosco_registro = QWidget()
        lay_reg_maestro = QVBoxLayout(self.p_kiosco_registro)
        lay_reg_maestro.setAlignment(Qt.AlignmentFlag.AlignCenter)

        container_form = QWidget()
        container_form.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        col_form = QVBoxLayout(container_form)
        col_form.setContentsMargins(0, 0, 0, 0)
        
        lay_header_reg = QHBoxLayout()
        lbl_reg_tit = QLabel("Formulario de Registro")
        lbl_reg_tit.setStyleSheet("font-size: 22px; font-weight: bold; color: #0F172A;")
        lbl_branding = self.cargar_imagen_branding(size=50)
        lay_header_reg.addWidget(lbl_reg_tit)
        lay_header_reg.addStretch()
        lay_header_reg.addWidget(lbl_branding)
        col_form.addLayout(lay_header_reg)
        
        form_reg = QFormLayout()
        form_reg.setSpacing(10)
        
        # --- NUEVO: Validadores ---
        self.validador_solo_numeros = QRegularExpressionValidator(QRegularExpression("^[0-9]*$"))
        self.validador_alfanumerico = QRegularExpressionValidator(QRegularExpression("^[a-zA-Z0-9]*$"))
        self.validador_solo_letras = QRegularExpressionValidator(QRegularExpression("^[a-zA-ZáéíóúÁÉÍÓÚñÑ\s]*$"))
        
        # --- Configuración de inputs ---
        estilo_input = "font-size: 16px; min-height: 40px; border: 2px solid #CBD5E1; border-radius: 6px; padding: 5px; background-color: white; color: #1E293B;"
        
        self.txt_reg_cedula = QLineEdit()
        self.txt_reg_cedula.setStyleSheet(estilo_input)
        self.txt_reg_cedula.setValidator(self.validador_solo_numeros) 
        
        self.txt_reg_matricula = QLineEdit()
        self.txt_reg_matricula.setStyleSheet(estilo_input)
        self.txt_reg_matricula.setValidator(self.validador_solo_numeros) 
        
        self.txt_reg_nombre = QLineEdit()
        self.txt_reg_nombre.setStyleSheet(estilo_input)
        self.txt_reg_nombre.setValidator(self.validador_solo_letras) 
        
        # --- Checkbox para Extranjero ---
        self.chk_extranjero = QCheckBox("Extranjero")
        self.chk_extranjero.setStyleSheet("font-size: 14px; font-weight: bold; color: #1E293B;")
        self.chk_extranjero.toggled.connect(self.alternar_modo_cedula)
        
        lay_cedula = QHBoxLayout()
        lay_cedula.addWidget(self.txt_reg_cedula)
        lay_cedula.addWidget(self.chk_extranjero)

        self.cmb_reg_tipo = QComboBox()
        self.cmb_reg_tipo.addItems(["Estudiante", "Profesor", "Empleado", "Invitado"])
        self.cmb_reg_tipo.setStyleSheet("background-color: white; color: #1E293B; font-size: 16px; min-height: 40px; border: 1px solid #CBD5E1;")
        self.cmb_reg_tipo.currentTextChanged.connect(self.evaluar_obligatoriedad_estudiante)
        
        self.cmb_reg_carrera = QComboBox()
        carreras_utesa = [
            "Ninguna / No Aplica", "Administración de Empresas", "Admin. de Emp. Turísticas y Hoteleras",
            "Arquitectura", "Comunicación Social", "Contabilidad", "Derecho", "Diseño de Interiores",
            "Educación", "Enfermería", "Ingeniería Civil", "Ingeniería Electrónica",
            "Ingeniería en Sistemas Computacionales", "Ingeniería Industrial", "Ingeniería Mecánica",
            "Lenguas Extranjeras/Modernas", "Marketing y Comunicación Digital", "Medicina",
            "Mercadeo", "Nutrición Humana y Dietética", "Odontología", "Optometría", "Psicología"
        ]
        self.cmb_reg_carrera.addItems(carreras_utesa)
        self.cmb_reg_carrera.setStyleSheet("background-color: white; color: #1E293B; font-size: 14px; min-height: 40px; border: 1px solid #CBD5E1;")
        
        self.lbl_alerta_estudiante = QLabel("* Matrícula y Carrera obligatorias para Estudiantes")
        self.lbl_alerta_estudiante.setStyleSheet("color: #EF4444; font-weight: bold; font-size: 12px;")
        
        form_reg.addRow(QLabel("Cédula o ID:"), lay_cedula)
        form_reg.addRow(QLabel("Matrícula:"), self.txt_reg_matricula)
        form_reg.addRow(QLabel("Nombre Completo:"), self.txt_reg_nombre)
        form_reg.addRow(QLabel("Tipo de Usuario:"), self.cmb_reg_tipo)
        form_reg.addRow(QLabel("Carrera:"), self.cmb_reg_carrera)
        form_reg.addRow("", self.lbl_alerta_estudiante)
        col_form.addLayout(form_reg)
        
        btn_reg_guardar = QPushButton("GUARDAR E INGRESAR")
        btn_reg_guardar.setStyleSheet("background-color: #10B981; color: white; font-size: 18px; font-weight: bold; min-height: 50px; border-radius: 8px;")
        btn_reg_guardar.clicked.connect(self.procesar_registro_usuario)
        col_form.addWidget(btn_reg_guardar)

        h_form_reg = QHBoxLayout()
        h_form_reg.addStretch(1)
        h_form_reg.addWidget(container_form, 4)
        h_form_reg.addStretch(1)

        # Teclado QWERTY
        col_teclado = QVBoxLayout()
        col_teclado.setAlignment(Qt.AlignmentFlag.AlignCenter)
        grid_qwerty = QGridLayout()
        grid_qwerty.setSpacing(5)
        grid_qwerty.setAlignment(Qt.AlignmentFlag.AlignCenter)
        
        filas_teclado = [
            ['1','2','3','4','5','6','7','8','9','0'],
            ['Q','W','E','R','T','Y','U','I','O','P'],
            ['A','S','D','F','G','H','J','K','L','Ñ'],
            ['Z','X','C','V','B','N','M','-','_','.'],
            ['Espacio', 'Borrar', 'Limpiar']
        ]
        
        for r_idx, fila in enumerate(filas_teclado):
            for c_idx, tecla in enumerate(fila):
                btn_tecla = QPushButton(tecla)
                btn_tecla.setFocusPolicy(Qt.FocusPolicy.NoFocus)
                btn_tecla.setMinimumHeight(55)
                btn_tecla.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
                btn_tecla.setStyleSheet("background-color: white; color: #1E293B; border: 1px solid #CBD5E1; font-size: 18px; font-weight: bold; border-radius: 6px;")
                if tecla == 'Espacio':
                    grid_qwerty.addWidget(btn_tecla, r_idx, c_idx, 1, 4) 
                elif tecla in ['Borrar', 'Limpiar']:
                    grid_qwerty.addWidget(btn_tecla, r_idx, c_idx + 3, 1, 3) 
                else:
                    grid_qwerty.addWidget(btn_tecla, r_idx, c_idx)
                btn_tecla.clicked.connect(self.tecla_alfa_presionada)
        
        col_teclado.addLayout(grid_qwerty)
        
        btn_reg_volver = QPushButton("Volver Atrás")
        btn_reg_volver.setStyleSheet("background-color: #64748B; color: white; font-size: 14px; min-height: 40px;")
        btn_reg_volver.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        btn_reg_volver.clicked.connect(self.limpiar_volver_kiosco)
        
        h_volver_reg = QHBoxLayout()
        h_volver_reg.addStretch(1)
        h_volver_reg.addWidget(btn_reg_volver, 3)
        h_volver_reg.addStretch(1)

        lay_reg_maestro.addStretch(1)
        lay_reg_maestro.addLayout(h_form_reg)
        lay_reg_maestro.addSpacing(20)
        lay_reg_maestro.addLayout(col_teclado)
        lay_reg_maestro.addSpacing(20)
        lay_reg_maestro.addLayout(h_volver_reg)
        lay_reg_maestro.addStretch(1)
        
        self.kiosco_stack.addWidget(self.p_kiosco_registro)
        layout_global.addLayout(layout_kiosco_master)

    def activar_modo_kiosco(self):
        self.stacked_widget.setCurrentIndex(1)
        self.kiosco_stack.setCurrentIndex(0)

    def alternar_pantalla_completa(self):
        if self.isFullScreen():
            dialogo = QDialog(self)
            dialogo.setWindowTitle("Seguridad")
            lay_d = QVBoxLayout(dialogo)
            lay_d.addWidget(QLabel("Contraseña administrativa para desbloquear:"))
            txt_pass = QLineEdit()
            txt_pass.setEchoMode(QLineEdit.EchoMode.Password)
            lay_d.addWidget(txt_pass)
            btn_d = QPushButton("Desbloquear y Restaurar Ventana")
            btn_d.setStyleSheet("background-color: #0F172A; color: white; min-height: 35px;")
            btn_d.clicked.connect(dialogo.accept)
            lay_d.addWidget(btn_d)
            
            if dialogo.exec() == QDialog.DialogCode.Accepted:
                if txt_pass.text() == "1234":
                    self.showMaximized()
                    self.nav_widget_kiosco.show()
                    self.btn_k_fullscreen.setText("🖵 BLOQUEAR EN PANTALLA COMPLETA")
                    self.btn_k_fullscreen.setStyleSheet("background-color: #1E293B; color: white; font-size: 16px; font-weight: bold; border-radius: 8px; padding: 15px;")
                else:
                    QMessageBox.critical(self, "Error", "Clave incorrecta")
        else:
            self.showFullScreen()
            self.nav_widget_kiosco.hide()
            self.btn_k_fullscreen.setText("🔓 DESBLOQUEAR Y SALIR DE PANTALLA COMPLETA")
            self.btn_k_fullscreen.setStyleSheet("background-color: #EF4444; color: white; font-size: 16px; font-weight: bold; border-radius: 8px; padding: 15px;")

    def verificar_cambio_biblioteca(self, index):
        if self.index_biblioteca_anterior != 0:
            dialogo = QDialog(self)
            dialogo.setWindowTitle("Seguridad - Cambio de Recinto")
            lay_d = QVBoxLayout(dialogo)
            lay_d.addWidget(QLabel("Contraseña administrativa para cambiar recinto:"))
            txt_pass = QLineEdit()
            txt_pass.setEchoMode(QLineEdit.EchoMode.Password)
            lay_d.addWidget(txt_pass)
            btn_d = QPushButton("Confirmar")
            btn_d.setStyleSheet("background-color: #0F172A; color: white; min-height: 35px;")
            btn_d.clicked.connect(dialogo.accept)
            lay_d.addWidget(btn_d)
            
            if dialogo.exec() == QDialog.DialogCode.Accepted and txt_pass.text() == "1234":
                self.index_biblioteca_anterior = index
            else:
                QMessageBox.warning(self, "Acceso Denegado", "Clave incorrecta. El recinto no se ha modificado.")
                self.cmb_kiosco_biblioteca.blockSignals(True)
                self.cmb_kiosco_biblioteca.setCurrentIndex(self.index_biblioteca_anterior)
                self.cmb_kiosco_biblioteca.blockSignals(False)
        else:
            self.index_biblioteca_anterior = index

    def alternar_modo_cedula(self, es_extranjero):
        if es_extranjero:
            self.txt_reg_cedula.setValidator(self.validador_alfanumerico)
        else:
            self.txt_reg_cedula.setValidator(self.validador_solo_numeros)
            texto_actual = self.txt_reg_cedula.text()
            if not texto_actual.isdigit():
                self.txt_reg_cedula.clear()

    def tecla_num_presionada(self):
        btn = self.sender()
        texto_tecla = btn.text()
        widget_activo = self.txt_ingreso_id
        
        texto_actual = widget_activo.text()
        if texto_tecla == "Limpiar": 
            widget_activo.clear()
            return
        elif texto_tecla == "Borrar": 
            widget_activo.setText(texto_actual[:-1])
            return
            
        nuevo_texto = texto_actual + texto_tecla
        
        validador = widget_activo.validator()
        if validador:
            estado, _, _ = validador.validate(nuevo_texto, len(nuevo_texto))
            if estado in [QValidator.State.Acceptable, QValidator.State.Intermediate]:
                widget_activo.setText(nuevo_texto)
        else:
            widget_activo.setText(nuevo_texto)

    def tecla_alfa_presionada(self):
        btn = self.sender()
        tecla = btn.text()
        
        widget_activo = QApplication.focusWidget()
        
        if isinstance(widget_activo, QLineEdit):
            texto_actual = widget_activo.text()
            if tecla == "Limpiar":
                widget_activo.clear()
                return
            elif tecla == "Borrar":
                widget_activo.setText(texto_actual[:-1])
                return
            
            caracter = " " if tecla == "Espacio" else tecla
            nuevo_texto = texto_actual + caracter
            
            validador = widget_activo.validator()
            if validador:
                estado, _, _ = validador.validate(nuevo_texto, len(nuevo_texto))
                if estado in [QValidator.State.Acceptable, QValidator.State.Intermediate]:
                    widget_activo.setText(nuevo_texto)
            else:
                widget_activo.setText(nuevo_texto)

    def evaluar_obligatoriedad_estudiante(self, tipo):
        if tipo == "Estudiante":
            self.lbl_alerta_estudiante.setText("* Matrícula y Carrera obligatorias")
            self.lbl_alerta_estudiante.setStyleSheet("color: #EF4444; font-weight: bold;")
        else:
            self.lbl_alerta_estudiante.setText("Matrícula y Carrera opcionales")
            self.lbl_alerta_estudiante.setStyleSheet("color: #64748B;")

    def limpiar_volver_kiosco(self):
        self.txt_ingreso_id.clear()
        self.txt_reg_cedula.clear()
        self.txt_reg_matricula.clear()
        self.txt_reg_nombre.clear()
        self.cmb_reg_tipo.setCurrentIndex(0)
        self.cmb_reg_carrera.setCurrentIndex(0)
        self.chk_extranjero.setChecked(False)
        self.kiosco_stack.setCurrentIndex(0)

    def procesar_ingreso_salon(self):
        biblioteca = self.cmb_kiosco_biblioteca.currentText()
        if self.cmb_kiosco_biblioteca.currentIndex() == 0:
            QMessageBox.warning(self, "Aviso", "Debe seleccionar el recinto actual en la parte superior antes de continuar.")
            return

        ident = self.txt_ingreso_id.text().strip()
        if not ident: return
        res = base_datos.buscar_usuario(ident)
        if res:
            esta_bloqueado = False
            if len(res) > 4: 
                esta_bloqueado = res[4]
                
            if esta_bloqueado:
                QMessageBox.warning(self, "Acceso Denegado", "Tiene el acceso bloqueado.")
                self.limpiar_volver_kiosco()
                return
                
            base_datos.registrar_acceso(res[0], biblioteca)
            QMessageBox.information(self, "Adelante", f"Bienvenido/a {res[1]}")
            self.limpiar_volver_kiosco()
        else:
            QMessageBox.warning(self, "Error", "ID no registrada. Vaya a Nuevo Usuario.")

    def procesar_registro_usuario(self):
        biblioteca = self.cmb_kiosco_biblioteca.currentText()
        if self.cmb_kiosco_biblioteca.currentIndex() == 0:
            QMessageBox.warning(self, "Aviso", "Debe seleccionar el recinto actual en la parte superior antes de continuar.")
            return

        cedula = self.txt_reg_cedula.text().strip()
        matricula = self.txt_reg_matricula.text().strip()
        nombre = self.txt_reg_nombre.text().strip()
        tipo = self.cmb_reg_tipo.currentText()
        
        carrera = self.cmb_reg_carrera.currentText()
        if carrera == "Seleccione una carrera... (Opcional)" or carrera == "Ninguna / No Aplica":
            carrera = ""
        
        if not cedula or not nombre: 
            QMessageBox.warning(self, "Aviso", "La Cédula y el Nombre son campos obligatorios.")
            return
            
        if tipo == "Estudiante" and (not matricula or not carrera):
            QMessageBox.warning(self, "Aviso", "Indique su Matrícula y Carrera (Obligatorio para Estudiantes).")
            return
            
        try:
            base_datos.registrar_nuevo_usuario(cedula, matricula, nombre, tipo, carrera, biblioteca)
            QMessageBox.information(self, "Éxito", f"¡Registro completado {nombre}!")
            self.limpiar_volver_kiosco()
        except psycopg2.IntegrityError:
            QMessageBox.critical(self, "Error", "Esa Cédula o Matrícula ya existe en el sistema.")

    # -----------------------------------------------------------------
    # VISTA 3: MÉTRICAS Y EXCEL
    # -----------------------------------------------------------------
    def crear_vista_metricas(self):
        layout_global = QVBoxLayout(self.vista_metricas)
        layout_global.addLayout(self.barra_navegacion(2))
        
        lbl_m_titulo = QLabel("Panel de Estadísticas")
        lbl_m_titulo.setStyleSheet("font-size: 18px; font-weight: bold; color: #0F172A; margin-bottom: 10px;")
        layout_global.addWidget(lbl_m_titulo)
        
        form_m = QFormLayout()
        self.cmb_periodos = QComboBox()
        self.cmb_periodos.addItems(["Día actual", "Semana", "Último mes", "Último trimestre", "Cuatrimestre", "Último año"])
        self.cmb_periodos.setMinimumHeight(35)
        form_m.addRow("Seleccionar Rango:", self.cmb_periodos)
        layout_global.addLayout(form_m)
        
        btn_generar_excel = QPushButton("Generar Estadística en Excel")
        btn_generar_excel.setStyleSheet("background-color: #1E293B; color: white; font-size: 16px; font-weight: bold; min-height: 55px; border-radius: 8px;")
        btn_generar_excel.clicked.connect(self.exportar_metricas_excel)
        layout_global.addWidget(btn_generar_excel)
        layout_global.addStretch()

    def calcular_fechas(self):
        p = self.cmb_periodos.currentText()
        hoy = datetime.now()
        fin = hoy.strftime("%Y-%m-%d 23:59:59")
        
        if p == "Día actual": ini = hoy
        elif p == "Semana": ini = hoy - timedelta(days=7)
        elif p == "Último mes": ini = hoy - timedelta(days=30)
        elif p == "Último trimestre": ini = hoy - timedelta(days=90)
        elif p == "Cuatrimestre": ini = hoy - timedelta(days=120)
        elif p == "Último año": ini = hoy - timedelta(days=365)
        else: ini = hoy
            
        return ini.strftime("%Y-%m-%d 00:00:00"), fin

    def exportar_metricas_excel(self):
        try:
            f_inicio, f_fin = self.calcular_fechas()
            headers, datos = base_datos.obtener_datos_metricas(f_inicio, f_fin)
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Métricas"
            
            f_head = Font(bold=True, color="FFFFFF")
            fill_head = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            
            for col_idx, text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=text)
                cell.font = f_head
                cell.fill = fill_head
                
            for r_idx, fila in enumerate(datos, 2):
                for c_idx, val in enumerate(fila, 1):
                    ws.cell(row=r_idx, column=c_idx, value=val)
                    
            for col in ws.columns:
                l_max = max(len(str(c.value or '')) for c in col)
                ws.column_dimensions[col[0].column_letter].width = max(l_max + 2, 10)
                
            nombre = f"Reporte_Biblioteca_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            wb.save(nombre)
            
            QMessageBox.information(self, "Excel Generado", f"El archivo se guardó correctamente como:\n{nombre}")

        except PermissionError:
            QMessageBox.critical(self, "Archivo Abierto", "No se pudo generar el Excel.\nParece que tienes un reporte anterior abierto. Cierra Excel e inténtalo de nuevo.")
        except Exception as error_general:
            mensaje_error = traceback.format_exc()
            print(mensaje_error)
            QMessageBox.critical(self, "Error Crítico", f"Ocurrió un error al generar el Excel:\n\n{str(error_general)}\n\nRevisa la terminal para más detalles.")

    # -----------------------------------------------------------------
    # VISTA 4: GESTIÓN DE USUARIOS
    # -----------------------------------------------------------------
    def crear_vista_usuarios(self):
        layout = QVBoxLayout(self.vista_usuarios)
        layout.addLayout(self.barra_navegacion(3))

        lbl_titulo = QLabel("Gestión Integral de Usuarios")
        lbl_titulo.setStyleSheet("font-size: 20px; font-weight: bold; color: #0F172A; margin-bottom: 10px;")
        layout.addWidget(lbl_titulo)

        self.txt_buscar_gestion = QLineEdit()
        self.txt_buscar_gestion.setPlaceholderText("🔍 Buscar usuario por cédula o matrícula...")
        self.txt_buscar_gestion.setStyleSheet("font-size: 16px; padding: 8px; border: 2px solid #CBD5E1; border-radius: 6px;")
        self.txt_buscar_gestion.textChanged.connect(self.actualizar_tabla_usuarios)
        layout.addWidget(self.txt_buscar_gestion)

        self.tabla_usuarios = QTableWidget(0, 6)
        self.tabla_usuarios.setHorizontalHeaderLabels(["Cédula", "Matrícula", "Nombre", "Tipo", "Carrera", "Acciones"])
        self.tabla_usuarios.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.tabla_usuarios.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.ResizeToContents)
        layout.addWidget(self.tabla_usuarios)
        
        self.actualizar_tabla_usuarios()

    def actualizar_tabla_usuarios(self):
        self.tabla_usuarios.setRowCount(0)
        filtro = self.txt_buscar_gestion.text().strip()
        
        try:
            usuarios = base_datos.obtener_usuarios_gestion(filtro)
            for row_idx, row in enumerate(usuarios):
                self.tabla_usuarios.insertRow(row_idx)
                
                self.tabla_usuarios.setItem(row_idx, 0, QTableWidgetItem(str(row[1] or "")))
                self.tabla_usuarios.setItem(row_idx, 1, QTableWidgetItem(str(row[2] or "")))
                self.tabla_usuarios.setItem(row_idx, 2, QTableWidgetItem(str(row[3] or "")))
                self.tabla_usuarios.setItem(row_idx, 3, QTableWidgetItem(str(row[4] or "")))
                self.tabla_usuarios.setItem(row_idx, 4, QTableWidgetItem(str(row[5] or "")))
                
                widget_acciones = QWidget()
                lay_acciones = QHBoxLayout(widget_acciones)
                lay_acciones.setContentsMargins(5, 2, 5, 2)
                
                btn_editar = QPushButton("Editar")
                btn_editar.setStyleSheet("background-color: #F59E0B; color: white; padding: 4px 8px; border-radius: 4px; font-weight: bold;")
                btn_editar.clicked.connect(lambda checked, datos=row: self.editar_usuario(datos))
                
                btn_borrar = QPushButton("Borrar")
                btn_borrar.setStyleSheet("background-color: #EF4444; color: white; padding: 4px 8px; border-radius: 4px; font-weight: bold;")
                btn_borrar.clicked.connect(lambda checked, uid=row[0]: self.borrar_usuario(uid))
                
                btn_bloquear = QPushButton()
                esta_bloqueado = row[6]
                if esta_bloqueado:
                    btn_bloquear.setText("Desbloquear")
                    btn_bloquear.setStyleSheet("background-color: #3B82F6; color: white; padding: 4px 8px; border-radius: 4px; font-weight: bold;")
                else:
                    btn_bloquear.setText("Bloquear")
                    btn_bloquear.setStyleSheet("background-color: #6B7280; color: white; padding: 4px 8px; border-radius: 4px; font-weight: bold;")
                
                btn_bloquear.clicked.connect(lambda checked, uid=row[0], estado=esta_bloqueado: self.bloquear_usuario(uid, estado))
                
                lay_acciones.addWidget(btn_editar)
                lay_acciones.addWidget(btn_borrar)
                lay_acciones.addWidget(btn_bloquear)
                
                self.tabla_usuarios.setCellWidget(row_idx, 5, widget_acciones)
        except Exception as e:
            print("Error cargando usuarios:", e)

    def editar_usuario(self, datos_usuario):
        dialogo = DialogoEditarUsuario(datos_usuario, self)
        if dialogo.exec() == QDialog.DialogCode.Accepted:
            nuevos_datos = dialogo.obtener_datos()
            try:
                base_datos.actualizar_usuario(datos_usuario[0], *nuevos_datos)
                QMessageBox.information(self, "Éxito", "Usuario actualizado correctamente.")
                self.actualizar_tabla_usuarios()
            except Exception as e:
                QMessageBox.critical(self, "Error", f"No se pudo actualizar el usuario.\nDetalle: {e}")

    def borrar_usuario(self, id_usuario):
        msg = QMessageBox(self)
        msg.setWindowTitle("Confirmar Eliminación")
        msg.setText("¿Deseas eliminar ese usuario?")
        msg.setIcon(QMessageBox.Icon.Warning)
        
        btn_aceptar = msg.addButton("Aceptar", QMessageBox.ButtonRole.AcceptRole)
        btn_cancelar = msg.addButton("Cancelar", QMessageBox.ButtonRole.RejectRole)
        
        btn_aceptar.setStyleSheet("background-color: red; color: white; padding: 5px 15px; font-weight: bold;")
        btn_cancelar.setStyleSheet("background-color: green; color: white; padding: 5px 15px; font-weight: bold;")
        
        msg.exec()
        
        if msg.clickedButton() == btn_aceptar:
            try:
                base_datos.eliminar_usuario(id_usuario)
                QMessageBox.information(self, "Eliminado", "El usuario ha sido borrado de la base de datos.")
                self.actualizar_tabla_usuarios()
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Hubo un problema al borrar el usuario.\nDetalle: {e}")

    def bloquear_usuario(self, id_usuario, estado_actual):
        try:
            base_datos.alternar_bloqueo_usuario(id_usuario, estado_actual)
            self.actualizar_tabla_usuarios()
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo cambiar el estado de bloqueo.\nDetalle: {e}")